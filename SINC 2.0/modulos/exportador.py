import pandas as pd


def exportar_excel(
    tabela,
    totais,
    economia,
    caminho_saida
):
    """
    Exporta o resultado da comparação
    para um arquivo Excel.
    """

    try:

        with pd.ExcelWriter(
            caminho_saida,
            engine="openpyxl"
        ) as writer:

            tabela.to_excel(
                writer,
                sheet_name="Comparação",
                index=False
            )

            resumo = []

            for fornecedor, total in totais.items():

                resumo.append(
                    {
                        "Fornecedor": fornecedor,
                        "Total Comprado (R$)": round(
                            float(total),
                            2
                        )
                    }
                )

            resumo.append(
                {
                    "Fornecedor": "ECONOMIA TOTAL",
                    "Total Comprado (R$)": round(
                        float(economia),
                        2
                    )
                }
            )

            df_resumo = pd.DataFrame(
                resumo
            )

            df_resumo.to_excel(
                writer,
                sheet_name="Resumo",
                index=False
            )

            try:

                from openpyxl.styles import (
                    Font,
                    PatternFill,
                    Alignment
                )

                ws = writer.sheets[
                    "Comparação"
                ]

                preenchimento_cabecalho = PatternFill(
                    fill_type="solid",
                    fgColor="0D47A1"
                )

                preenchimento_menor = PatternFill(
                    fill_type="solid",
                    fgColor="C8F7C5"
                )

                fonte_cabecalho = Font(
                    color="FFFFFF",
                    bold=True
                )

                fonte_menor = Font(
                    color="1B5E20",
                    bold=True
                )

                # =================================================
                # CABEÇALHO
                # =================================================

                for celula in ws[1]:

                    celula.fill = (
                        preenchimento_cabecalho
                    )

                    celula.font = (
                        fonte_cabecalho
                    )

                    celula.alignment = Alignment(
                        horizontal="center"
                    )

                # =================================================
                # IDENTIFICA AS COLUNAS DE PREÇO
                # =================================================

                colunas_preco = []

                colunas_que_nao_sao_preco = {
                    "Código",
                    "Referência",
                    "Descrição",
                    "Marca",
                    "Qtd",
                    "Economia"
                }

                for coluna in range(
                    1,
                    ws.max_column + 1
                ):

                    nome_coluna = ws.cell(
                        row=1,
                        column=coluna
                    ).value

                    if not nome_coluna:
                        continue

                    nome_coluna = str(
                        nome_coluna
                    ).strip()

                    # Marca Cotada não é preço.

                    if (
                        " Marca Cotada"
                        in nome_coluna
                    ):
                        continue

                    # Colunas fixas não são preço.

                    if (
                        nome_coluna
                        in colunas_que_nao_sao_preco
                    ):
                        continue

                    colunas_preco.append(
                        coluna
                    )

                # =================================================
                # PINTA O MENOR PREÇO DE CADA ITEM
                # =================================================

                for linha in range(
                    2,
                    ws.max_row + 1
                ):

                    precos = []

                    for coluna in colunas_preco:

                        celula = ws.cell(
                            row=linha,
                            column=coluna
                        )

                        valor = celula.value

                        if valor is None:
                            continue

                        try:

                            valor_numerico = float(
                                valor
                            )

                        except (
                            ValueError,
                            TypeError
                        ):

                            continue

                        if valor_numerico <= 0:
                            continue

                        precos.append(
                            (
                                coluna,
                                valor_numerico
                            )
                        )

                    if not precos:
                        continue

                    menor_preco = min(
                        preco
                        for coluna, preco
                        in precos
                    )

                    # ---------------------------------------------
                    # PINTA TODOS OS EMPATES
                    # ---------------------------------------------

                    for coluna, preco in precos:

                        if abs(
                            preco - menor_preco
                        ) < 0.000001:

                            celula = ws.cell(
                                row=linha,
                                column=coluna
                            )

                            celula.fill = (
                                preenchimento_menor
                            )

                            celula.font = (
                                fonte_menor
                            )

                # =================================================
                # AJUSTA LARGURA DAS COLUNAS
                # =================================================

                for coluna in ws.columns:

                    maior = 0

                    letra = coluna[0].column_letter

                    for celula in coluna:

                        if celula.value is not None:

                            tamanho = len(
                                str(
                                    celula.value
                                )
                            )

                            if tamanho > maior:
                                maior = tamanho

                    ws.column_dimensions[
                        letra
                    ].width = min(
                        maior + 2,
                        50
                    )

                ws.freeze_panes = "A2"

                # =================================================
                # RESUMO
                # =================================================

                ws_resumo = writer.sheets[
                    "Resumo"
                ]

                for celula in ws_resumo[1]:

                    celula.fill = (
                        preenchimento_cabecalho
                    )

                    celula.font = (
                        fonte_cabecalho
                    )

                    celula.alignment = Alignment(
                        horizontal="center"
                    )

                ws_resumo.column_dimensions[
                    "A"
                ].width = 35

                ws_resumo.column_dimensions[
                    "B"
                ].width = 25

                for linha in range(
                    2,
                    ws_resumo.max_row + 1
                ):

                    ws_resumo.cell(
                        row=linha,
                        column=2
                    ).number_format = (
                        'R$ #,##0.00'
                    )

                ws_resumo.freeze_panes = "A2"

            except Exception:
                pass

    except Exception as erro:

        raise Exception(
            f"ERRO AO EXPORTAR EXCEL: {erro}"
        )